from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from datetime import datetime

def generar_excel_herramientas(herramientas, usuario_nombre, usuario_rol):
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario Herramientas"

    # Estilos
    header_font = Font(name='Arial', size=14, bold=True, color="2196F3")
    info_font = Font(name='Arial', size=10, color="505050")
    table_header_font = Font(name='Arial', size=11, bold=True, color="FFFFFF")
    table_header_fill = PatternFill(start_color="2196F3", end_color="2196F3", fill_type="solid")
    center_alignment = Alignment(horizontal='center', vertical='center')
    left_alignment = Alignment(horizontal='left', vertical='center')
    
    border_side = Side(style='thin', color="B0B0B0")
    table_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    # --- ENCABEZADO ---
    ws.merge_cells('A1:F1')
    ws['A1'] = "CELUCENTRO - REPORTE DE INVENTARIO"
    ws['A1'].font = header_font
    ws['A1'].alignment = Alignment(horizontal='center')

    ws.merge_cells('A2:F2')
    ws['A2'] = f"Generado por: {usuario_nombre} ({usuario_rol})"
    ws['A2'].font = info_font
    ws['A2'].alignment = Alignment(horizontal='right')

    ws.merge_cells('A3:F3')
    ws['A3'] = f"Fecha de Emisión: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws['A3'].font = info_font
    ws['A3'].alignment = Alignment(horizontal='right')

    # --- RESUMEN ---
    total = herramientas.count()
    bajo_stock = herramientas.filter(stock_disponible__gt=0, stock_disponible__lte=3).count()
    agotados = herramientas.filter(stock_disponible__lte=0).count()

    ws['A5'] = "Resumen de Inventario"
    ws['A5'].font = Font(bold=True, size=12)
    
    ws['A6'] = f"Total Herramientas: {total}"
    ws['B6'] = f"Con Stock Bajo: {bajo_stock}"
    ws['C6'] = f"Agotadas: {agotados}"
    for cell in ['A6', 'B6', 'C6']:
        ws[cell].font = info_font

    # --- TABLA ---
    headers = ['#', 'Nombre de Herramienta', 'Categoría', 'Stock Total', 'Stock Disponible', 'Estado']
    start_row = 8
    
    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_num)
        cell.value = header_title
        cell.font = table_header_font
        cell.fill = table_header_fill
        cell.alignment = center_alignment
        cell.border = table_border

    # Datos
    for row_num, (i, h) in enumerate(enumerate(herramientas, 1), start_row + 1):
        # Numeración
        ws.cell(row=row_num, column=1, value=i).alignment = center_alignment
        
        # Nombre
        ws.cell(row=row_num, column=2, value=h.nombre).alignment = left_alignment
        
        # Categoría
        cat = h.id_categoria.nombre if h.id_categoria else 'General'
        ws.cell(row=row_num, column=3, value=cat).alignment = left_alignment
        
        # Stock Total
        ws.cell(row=row_num, column=4, value=h.stock_total).alignment = center_alignment
        
        # Stock Disponible
        disp_cell = ws.cell(row=row_num, column=5, value=h.stock_disponible)
        disp_cell.alignment = center_alignment
        if h.stock_disponible <= 0:
            disp_cell.font = Font(color="FF0000", bold=True) # Rojo
        elif h.stock_disponible <= 3:
            disp_cell.font = Font(color="FF8C00", bold=True) # Naranja

        # Estado
        estado = 'ACTIVO' if h.estado else 'INACTIVO'
        ws.cell(row=row_num, column=6, value=estado).alignment = center_alignment

        # Bordes para toda la fila
        for col in range(1, 7):
            ws.cell(row=row_num, column=col).border = table_border

    # Ajustar anchos de columna
    dims = {}
    for row in ws.rows:
        for cell in row:
            if cell.value:
                dims[cell.column_letter] = max((dims.get(cell.column_letter, 0), len(str(cell.value))))
    for col, value in dims.items():
        ws.column_dimensions[col].width = value + 5

    return wb
